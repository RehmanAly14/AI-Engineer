import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel
load_dotenv()

my_api_key = os.getenv("Groq_API_KEY")
if not my_api_key:
    raise ValueError("Groq_API_KEY not found in environment variables.")

client = Groq(api_key=my_api_key)

model = "llama-3.3-70b-versatile"

job_description =""""
Description
Amazon Software Development Engineer (SDE) roles focus on building highly scalable, fault-tolerant distributed systems and cloud-native architectures. Key responsibilities include writing clean code, participating in CI/CD practices, and managing the full software lifecycle—from initial design to ongoing production operations and on-call support.A standard Amazon job posting typically includes the following components:Key ResponsibilitiesSystem Design & Development: Architect and build large-scale microservices.Operational Excellence: Maintain and scale resilient, distributed systems, and participate in on-call rotations to troubleshoot production issues.Collaboration: Partner closely with UX designers, product managers, and cross-disciplinary teams to deliver customer-focused solutions.Innovation: Implement modern development practices and leverage AI-powered tools to enhance productivity.Basic QualificationsExperience: 1+ years of non-internship professional software development experience (for SDE-I), or specific enrollment/recent graduation for University postings.Coding Skills: Proficiency in at least one modern programming language such as Java, Python, C++, C#, Go, or Rust.CS Fundamentals: Strong understanding of data structures, algorithms, object-oriented design, and complexity analysis.Preferred QualificationsExperience building or maintaining cloud-native applications, ideally within distributed, multi-tiered systems.Strong track record of owning the full software development lifecycle and handling ambiguity.Excellent verbal and written communication skills
"""

class Job (BaseModel):
   role: str
   required_skills: list[str]
   preferred_skills: list[str]
   minimum_experience: float | None
   education_requirements: list[str] 
   responsibilities: list[str]

job_schema = Job.model_json_schema()

response_format = {
    'type': "json_object",
}

system_prompt = f"""
You are an expert HR recruiter.

Extract information from the following job description.

Return ONLY valid JSON.

The JSON must match this schema:

{job_schema}

Do not return the schema itself.
Do not return fields like "properties", "title", or "type".

Fill the schema with extracted information only.
If a field is missing, use null or [].
Do not invent any information.
"""

user_prompt = f"""
Extract the job information from the following description.

Return ONLY JSON.

Job Description:

{job_description}
"""

    

message_system ={
    "role": "system",
    "content": system_prompt
}
message = {
    "role": "user",
    "content": user_prompt
}
message_list = [message_system, message]
response = client.chat.completions.create(model=model, messages=message_list,response_format=response_format)
answer=response.choices[0].message.content





import json
raw_json = answer
data_file = json.loads(raw_json)
job = Job(**data_file)


print(job.role)
print(job.required_skills)
print(job.preferred_skills)





# part 2 define resume schema and extract information from resume

class matchResult (BaseModel):
    score: float
    details: dict

class Experience (BaseModel):
    company: str | None=None
    role: str | None=None
    duration: str | None=None
    responsibilities: list[str] | None=None
    description: str | None=None
    skills_used: list[str] | None=None

class Resume (BaseModel):
    name: str | None=None
    email: str | None=None
    phone: str | None=None
    total_experience_years: float | None=None
    education: list[str] = []
    skills: list[str] = []
    experiences: list[Experience] = []
    projects: list[str] = []
    certifications: list[str] = []

resume_schema = Resume.model_json_schema()

def final_score(job,resume):
    match_schema = matchResult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:

    {match_schema}

    Give me:

    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """
    message={
        "role": "user",
        "content" : prompt
    }
    messages=[message]
    response_format={
        "type": "json_object"
    }
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return matchResult(**data)

def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume

# part 3 pdf parsing and extracting information from resume

from pypdf import PdfReader
from docx import Document

def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"

    return text

def read_docx(file_path):
    doc = Document(file_path)
    text = ""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else: 
        return None
    

# let do it now
import time
resume_folder = Path("resumes")
all_results=[]
for file_path in resume_folder.iterdir():
    #C:\Users\Pratyush\padho_with_pratyush\week1\day5\resumes\abhay resume new - Abhay Singh.pdf
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue
    print("\nProcessing:", file_path.name)
    resume_text = read_resume(file_path)
    parsed_resume=parse_resume(resume_text) # llm call1
    time.sleep(5)
    result = final_score(job, parsed_resume) #llm caLL2
    #score and details
    #acount chtgpt
    # request bhejna shhur krega millions
    #chattgot server jam ho jayega
    time.sleep(5)
    print("Score:", result.score)
    all_results.append({
        "name": parsed_resume.name,
        "score": result.score,
        "details": result.details
    })
all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)
top_2 = all_results[:2]
worst_2 = all_results[-2:]


print("TOP 2 CANDIDATES")
for candidate in top_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )

    print(candidate["details"])

print("LOWEST 2 CANDIDATES")
for candidate in worst_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])